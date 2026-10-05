"""
SRS Parser: extracts section hierarchy and requirements from DOCX and PDF.
"""
import re
import json
from typing import List, Dict, Optional, Tuple

def _is_section_heading(text: str) -> bool:
    """Detect lines like '4.3 Something' or '1. Introduction'."""
    return bool(re.match(r'^\d+(\.\d+)*\.?\s+\S', text.strip()))

def _extract_req_id(text: str) -> Optional[str]:
    """Find requirement identifiers like FR-101, NFR-05, REQ-001."""
    m = re.search(r'\b(FR|NFR|REQ|UC|R)-?\d+\b', text, re.IGNORECASE)
    return m.group(0).upper() if m else None

def _starts_with_req_id(text: str) -> bool:
    return bool(re.match(r'^(FR|NFR|REQ|UC|R)-?\d+\b', text.strip(), re.IGNORECASE))

def _has_shall(text: str) -> bool:
    return 'shall' in text.lower()

def _classify_requirement(text: str, section: str) -> str:
    nfr_keywords = ['performance', 'security', 'usability', 'reliability', 'maintainability',
                    'portability', 'compatibility', 'availability', 'scalab', 'cost']
    section_lower = section.lower()
    if 'non-functional' in section_lower or any(k in section_lower for k in nfr_keywords):
        return 'non_functional'
    if any(k in text.lower() for k in nfr_keywords):
        return 'non_functional'
    return 'functional'

def extract_from_docx(file_path: str) -> Dict:
    from docx import Document
    doc = Document(file_path)
    sections = []
    requirements = []
    current_section = "Unknown"
    
    current_req = None

    def process_text(text: str, source_ref: str):
        nonlocal current_section, current_req
        if not text:
            return
            
        if _is_section_heading(text):
            current_section = text
            sections.append(text)
            current_req = None
            return

        is_new_req = _has_shall(text) or _starts_with_req_id(text)
        
        if is_new_req:
            req_id = _extract_req_id(text)
            req_class = _classify_requirement(text, current_section)
            current_req = {
                "id": req_id,
                "text": text,
                "section": current_section,
                "class": req_class,
                "source_ref": source_ref
            }
            requirements.append(current_req)
        elif current_req is not None:
            # Join wrapped lines
            current_req["text"] += " " + text

    # Process paragraphs
    for i, para in enumerate(doc.paragraphs):
        process_text(para.text.strip(), f"Paragraph {i+1}")
        
    # Process tables
    for i, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                process_text(cell.text.strip(), f"Table {i+1}, Row {r_idx+1}, Col {c_idx+1}")

    return {"sections": sections, "requirements": requirements}

def extract_from_pdf(file_path: str) -> Dict:
    import pdfplumber
    sections = []
    requirements = []
    current_section = "Unknown"
    current_req = None

    with pdfplumber.open(file_path) as pdf:
        for p_idx, page in enumerate(pdf.pages):
            text = page.extract_text() or ""
            for line in text.split('\n'):
                line = line.strip()
                if not line:
                    continue
                    
                if _is_section_heading(line):
                    current_section = line
                    sections.append(line)
                    current_req = None
                    continue

                is_new_req = _has_shall(line) or _starts_with_req_id(line)
                
                if is_new_req:
                    req_id = _extract_req_id(line)
                    req_class = _classify_requirement(line, current_section)
                    current_req = {
                        "id": req_id,
                        "text": line,
                        "section": current_section,
                        "class": req_class,
                        "source_ref": f"Page {p_idx+1}"
                    }
                    requirements.append(current_req)
                elif current_req is not None:
                    # Join wrapped lines
                    current_req["text"] += " " + line

    return {"sections": sections, "requirements": requirements}

def parse_srs(file_path: str, file_format: str) -> Dict:
    if file_format == "pdf":
        return extract_from_pdf(file_path)
    return extract_from_docx(file_path)
