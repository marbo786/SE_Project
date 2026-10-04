"""
Seed script: populates the database with demo users, assignments, and submissions.
Usage:
    python backend/scripts/seed_demo.py
"""
import os
import sys

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))
os.chdir(os.path.abspath(os.path.join(os.path.dirname(__file__), '..')))

from app.core.database import SessionLocal, engine, Base
from app.core.security import get_password_hash
from app.models.user import User
from app.models.assignment import Assignment
from app.models.submission import Submission
from app.models.artifact import Artifact
from app.pipeline.analyzer import run_analysis_pipeline

def seed():
    Base.metadata.create_all(bind=engine)
    db = SessionLocal()

    print("Creating demo accounts...")
    # 1. Instructor
    instructor = db.query(User).filter(User.email == "israr@giki.edu.pk").first()
    if not instructor:
        instructor = User(
            name="Israr Ahmad",
            email="israr@giki.edu.pk",
            hashed_password=get_password_hash("instructor123"),
            role="instructor"
        )
        db.add(instructor)
        db.commit()
        db.refresh(instructor)
        print("âœ“ Created Instructor: israr@giki.edu.pk / instructor123")
    else:
        print("  Instructor already exists")

    # 2. Student
    student = db.query(User).filter(User.email == "mohsin@student.giki.edu.pk").first()
    if not student:
        student = User(
            name="Mohsin Saeed",
            email="mohsin@student.giki.edu.pk",
            hashed_password=get_password_hash("student123"),
            role="student"
        )
        db.add(student)
        db.commit()
        db.refresh(student)
        print("âœ“ Created Student: mohsin@student.giki.edu.pk / student123")
    else:
        print("  Student already exists")

    # 3. Assignment
    assignment = db.query(Assignment).filter(Assignment.course_code == "CS325").first()
    if not assignment:
        assignment = Assignment(
            title="CS325 Term Project: Software Requirements & UML Model",
            course_code="CS325",
            instructor_id=instructor.id
        )
        db.add(assignment)
        db.commit()
        db.refresh(assignment)
        print("âœ“ Created Assignment: CS325 Term Project")
    else:
        print("  Assignment already exists")

    # 4. Demo Submission with realistic SRS & UML
    submission = db.query(Submission).filter(Submission.team_name == "Team Nova").first()
    if not submission:
        submission = Submission(
            assignment_id=assignment.id,
            student_id=student.id,
            team_name="Team Nova",
            member_names='["Mohsin Saeed", "Hamza Sami", "Ahmad Sajid", "M. Taha"]',
            version=1,
            status="pending"
        )
        db.add(submission)
        db.commit()
        db.refresh(submission)
        print(f"âœ“ Created Submission for Team Nova (ID: {submission.id})")

        # Create demo SRS docx
        import io
        from docx import Document
        doc = Document()
        doc.add_heading("1. Introduction", level=1)
        doc.add_paragraph("This SRS defines the requirements for an Autonomous Drone Dispatching System.")
        doc.add_heading("4. Functional requirements", level=1)
        doc.add_paragraph("FR-101: The drone controller shall provide fast, reliable, and user-friendly fleet telemetry.")
        doc.add_paragraph("FR-102: The pilot shall dispatch drones and the pilot shall abort missions when weather is severe.")
        doc.add_paragraph("FR-103: The flight coordinator shall record telemetry metrics every 500ms.")
        doc.add_heading("5. Non-functional requirements", level=1)
        doc.add_paragraph("NFR-01: The system shall be fast and flexible.")
        doc.add_paragraph("NFR-02: Telemetry latency shall not exceed 250ms.")

        uploads_dir = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "..", "uploads", str(submission.id)))
        os.makedirs(uploads_dir, exist_ok=True)
        srs_path = os.path.join(uploads_dir, "drone_srs.docx")
        doc.save(srs_path)

        srs_artifact = Artifact(
            submission_id=submission.id,
            artifact_type="srs",
            file_format="docx",
            file_path=srs_path,
            original_filename="drone_srs.docx"
        )
        db.add(srs_artifact)

        # PlantUML Use Case
        uml_usecase_path = os.path.join(uploads_dir, "usecase.puml")
        with open(uml_usecase_path, "w", encoding="utf-8") as f:
            f.write("""@startuml
left to right direction
actor "Drone Pilot" as Pilot
actor "Maintenance Tech" as Tech
(Dispatch Drone)
(Abort Mission)
(Record Telemetry)
(Orphaned Use Case)
Pilot --> (Dispatch Drone)
Pilot --> (Abort Mission)
Pilot --> (Record Telemetry)
@enduml""")

        uml_artifact = Artifact(
            submission_id=submission.id,
            artifact_type="uml_usecase",
            file_format="plantuml",
            file_path=uml_usecase_path,
            original_filename="usecase.puml"
        )
        db.add(uml_artifact)
        db.commit()

        # Run analysis pipeline
        print("Running analysis pipeline on demo submission...")
        run_analysis_pipeline(submission.id)
        print("âœ“ Analysis complete! Scores and findings generated.")
    else:
        print("  Demo submission already exists")

    db.close()
    print("\nâœ… Seed completed successfully! You can login as:")
    print("   Instructor: israr@giki.edu.pk / instructor123")
    print("   Student:    mohsin@student.giki.edu.pk / student123")

if __name__ == "__main__":
    seed()
